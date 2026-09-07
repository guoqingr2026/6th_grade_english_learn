# -*- coding: utf-8 -*-
"""Learning Companion V2.0 — reference layout: vocab page + sentence page + 3-module sidebar."""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from unit_data import UNITS, UNIT_INCLUDES, MOTIVATIONAL_QUOTES

OUTPUT = r"C:\Users\guoqren\.cursor\englishi learn\Learning Companion V2.0 六年级上册 Unit1-6.docx"

BLACK = RGBColor(0x00, 0x00, 0x00)
BLUE = RGBColor(0x7C, 0xB8, 0xE8)
ORANGE = RGBColor(0xED, 0x7D, 0x31)
GRAY = RGBColor(0x80, 0x80, 0x80)
DARK = RGBColor(0x33, 0x33, 0x33)

LEARNING_BLUE = "7CB8E8"
LEARNING_BLUE_LIGHT = "B8D9F2"
HEX_ORANGE = "ED7D31"
HEX_LINE = "D0D0D0"

FONT_CN = "微软雅黑"
FONT_EN = "Calibri"
OTWSRR = "Observe → Think → Write → Speak → Recall → Reflect"

FOUR_LINE_SPECS = [
    (0.34, LEARNING_BLUE, "single", "8"),
    (0.32, LEARNING_BLUE_LIGHT, "dashed", "4"),
    (0.32, LEARNING_BLUE_LIGHT, "dashed", "4"),
    (0.34, LEARNING_BLUE, "single", "10"),
]

FOUR_LINE_COMPACT = [
    (0.26, LEARNING_BLUE, "single", "6"),
    (0.24, LEARNING_BLUE_LIGHT, "dashed", "4"),
    (0.24, LEARNING_BLUE_LIGHT, "dashed", "4"),
    (0.26, LEARNING_BLUE, "single", "8"),
]

FOUR_LINE_COMPACT_H = sum(h for h, _, _, _ in FOUR_LINE_COMPACT)
FOUR_LINE_GRID_H = sum(h for h, _, _, _ in FOUR_LINE_SPECS)

# 左侧栏可书写高度（cm），用于手动分页并保留右侧三模块笔记
LEFT_COL_BUDGET_CM = 18.2

TEXT_LINE_SPACING = Pt(15)
MODULE_SPACING_PT = 5
TEXT_GRID_GAP_CM = 0.12
MODULE_TAIL_CM = 0.10

# Sidebar — must fit one A4 with ruled lines (total ~15.5cm boxes + labels)
SB_NOTES_H = 7.2
SB_NOTES_LINES = 10
SB_POINTS_H = 3.6
SB_POINTS_LINES = 5
SB_DISCOVERY_H = 3.6
SB_DISCOVERY_LINES = 5

PAGE_LABELS = [
    "Unit Guide · 学习导航",
    "Vocabulary · 核心词汇",
    "Phrases · 重点短语",
    "Sentences · 重点句型",
    "Vocab Recall · 词汇默写",
    "Sentence Dictation · 句型默写",
    "Error Log · 易错整理",
    "Part B Learning · 课堂学习",
    "Reading Companion · 阅读记录",
    "Grammar Discovery · 语法发现",
    "Knowledge Map · 思维导图",
    "Mini Challenge · 综合挑战",
    "Growth Journal · 成长记录",
]


# ── helpers ──────────────────────────────────────────────────────────

def set_run_font(run, size=10, bold=False, color=None, italic=False):
    run.font.name = FONT_EN
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), FONT_EN)
    rFonts.set(qn("w:hAnsi"), FONT_EN)
    rFonts.set(qn("w:eastAsia"), FONT_CN)
    rPr.insert(0, rFonts)


def add_text(paragraph, text, size=10, bold=False, color=BLACK, italic=False, label=False):
    run = paragraph.add_run(text)
    set_run_font(run, size, bold=(bold or label), color=color, italic=italic)
    return run


def set_cell_shading(cell, color_hex):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color_hex)
    shading.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_border(cell, edge, val="single", sz="10", color=LEARNING_BLUE):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), val)
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), "0")
    el.set(qn("w:color"), color)
    tcBorders.append(el)


def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    for edge, spec in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        if spec:
            set_cell_border(cell, edge, spec.get("val", "single"),
                            spec.get("sz", "10"), spec.get("color", LEARNING_BLUE))


def set_row_height(row, cm, rule="exact"):
    trPr = row._tr.get_or_add_trPr()
    for old in trPr.findall(qn("w:trHeight")):
        trPr.remove(old)
    trHeight = OxmlElement("w:trHeight")
    trHeight.set(qn("w:val"), str(int(cm * 567)))
    trHeight.set(qn("w:hRule"), rule)
    trPr.append(trHeight)


def set_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    if trPr.find(qn("w:cantSplit")) is None:
        trPr.append(OxmlElement("w:cantSplit"))


def set_table_width(table, cm):
    tblPr = table._tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        table._tbl.insert(0, tblPr)
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), str(int(cm * 567)))
    tblW.set(qn("w:type"), "dxa")
    tblPr.append(tblW)


def remove_table_borders(table):
    tblPr = table._tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        table._tbl.insert(0, tblPr)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)


def learning_frame():
    b = {"val": "single", "sz": "10", "color": LEARNING_BLUE}
    return dict(top=b, bottom=b, left=b, right=b)


def cell_text(cell, text, size=8, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = Pt(max(size + 4, 11))
    add_text(p, text, size=size, bold=bold, label=bold, color=BLACK)


def set_cell_margins(cell, top=0, bottom=0, left=0, right=0):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for edge, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(int(val)))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


def estimate_lines(text, chars_per_line=30):
    if not text:
        return 1
    return max(1, (len(text) + chars_per_line - 1) // chars_per_line)


def add_label(parent, text, size=8, tight=False):
    p = parent.add_paragraph()
    p.paragraph_format.space_before = Pt(2 if tight else 3)
    p.paragraph_format.space_after = Pt(1)
    add_text(p, text, size=size, label=True)




def add_ruled_box(parent, height_cm, width_cm=6.5, lines=8, dashed=False):
    """蓝色学习框 + 内部横线格."""
    table = parent.add_table(rows=1, cols=1)
    remove_table_borders(table)
    set_table_width(table, width_cm)
    outer = table.rows[0].cells[0]
    set_cell_shading(outer, "FFFFFF")
    set_cell_borders(outer, **learning_frame())
    set_row_height(table.rows[0], height_cm)
    inner = outer.add_table(rows=lines, cols=1)
    remove_table_borders(inner)
    rh = max(0.32, (height_cm - 0.15) / lines)
    border_val = "dashed" if dashed else "single"
    for row in inner.rows:
        set_row_height(row, rh)
        c = row.cells[0]
        set_cell_borders(c, bottom={"val": border_val, "sz": "4", "color": HEX_LINE})
        c.paragraphs[0].paragraph_format.space_before = Pt(0)
        c.paragraphs[0].paragraph_format.space_after = Pt(0)
    return outer


def add_learning_box(parent, height_cm, width_cm=6.5, ruled_lines=0):
    if ruled_lines > 0:
        return add_ruled_box(parent, height_cm, width_cm, lines=ruled_lines)
    table = parent.add_table(rows=1, cols=1)
    remove_table_borders(table)
    set_table_width(table, width_cm)
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, "FFFFFF")
    set_cell_borders(cell, **learning_frame())
    set_row_height(table.rows[0], height_cm)
    return cell


def apply_grid_line_row(row, h, color, val, sz, width_cm=10.5):
    """在表格的单独一行上绘制一条四线格线."""
    set_row_height(row, h)
    set_row_cant_split(row)
    c = row.cells[0]
    set_cell_margins(c, top=0, bottom=0, left=0, right=0)
    set_cell_borders(c, bottom={"val": val, "sz": sz, "color": color})
    para = c.paragraphs[0]
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(0)
    para.paragraph_format.line_spacing = Pt(1)


def add_four_line_grid(parent, width_cm=10.5, specs=None):
    """标准四线格模块 — 4 行固定高度，每行一条线."""
    specs = specs or FOUR_LINE_SPECS
    table = parent.add_table(rows=4, cols=1)
    remove_table_borders(table)
    set_table_width(table, width_cm)
    for i, spec in enumerate(specs):
        apply_grid_line_row(table.rows[i], *spec, width_cm=width_cm)
    return table


def inline_text_height(num, en, cn, chars_per_line=36):
    prefix = f"{num}. " if num else ""
    combined = f"{prefix}{en}  {cn}" if cn else f"{prefix}{en}"
    lines = estimate_lines(combined, chars_per_line)
    return max(0.55, 0.40 * lines) + TEXT_GRID_GAP_CM


def add_inline_en_cn(paragraph, num, en, cn):
    """英文与中文同一行，过长自动换行，不裁切."""
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = TEXT_LINE_SPACING
    prefix = f"{num}. " if num else ""
    add_text(paragraph, f"{prefix}{en}", size=8, label=True)
    if cn:
        add_text(paragraph, "  ", size=8)
        add_text(paragraph, cn, size=8, color=BLACK)


def add_module_spacer(parent):
    p = parent.add_paragraph()
    p.paragraph_format.space_before = Pt(MODULE_SPACING_PT)
    p.paragraph_format.space_after = Pt(MODULE_SPACING_PT)
    p.paragraph_format.line_spacing = Pt(1)


def phrase_text_height(en, cn, num=1):
    return inline_text_height(num, en, cn)


def sentence_text_height(en, cn, num=1):
    return inline_text_height(num, en, cn, chars_per_line=34)


def add_vocab_table(parent, vocab_list, width_cm=10.5, compact=False):
    """词汇表 — compact 模式 8 号字，按词数自动压缩行高与列宽."""
    headers = ["单词 Word", "音标 Phonetics", "中文释义 Meaning"]
    n = len(vocab_list)
    table = parent.add_table(rows=1 + n, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table, width_cm)
    fs = 8 if compact else 7

    if compact:
        if n <= 12:
            row_h, col_ratios = 0.52, (0.36, 0.28, 0.36)
        elif n <= 16:
            row_h, col_ratios = 0.48, (0.35, 0.27, 0.38)
        else:
            row_h, col_ratios = 0.45, (0.34, 0.26, 0.40)
    else:
        row_h, col_ratios = 0.54, (0.36, 0.28, 0.36)

    set_row_height(table.rows[0], row_h + 0.06)

    for i, h in enumerate(headers):
        c = table.rows[0].cells[i]
        set_cell_shading(c, "FFFFFF")
        set_cell_borders(c, **learning_frame())
        cell_text(c, h, size=fs, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        c.width = Cm(width_cm * col_ratios[i])

    for ri, (word, ipa, pos, cn) in enumerate(vocab_list, 1):
        row = table.rows[ri]
        word_lines = estimate_lines(f"{ri}. {word}", 22)
        cn_lines = estimate_lines(f"{cn} ({pos})", 18)
        rh = max(row_h, 0.38 + 0.34 * max(word_lines, cn_lines))
        set_row_height(row, rh)
        vals = [f"{ri}. {word}", ipa, f"{cn}  ({pos})"]
        aligns = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT]
        for ci, (val, align) in enumerate(zip(vals, aligns)):
            c = row.cells[ci]
            set_cell_shading(c, "FFFFFF")
            set_cell_margins(c, top=40, bottom=40, left=60, right=60)
            set_cell_borders(c,
                             bottom={"val": "single", "sz": "4", "color": LEARNING_BLUE},
                             left={"val": "single", "sz": "4", "color": LEARNING_BLUE},
                             right={"val": "single", "sz": "4", "color": LEARNING_BLUE})
            cell_text(c, val, size=fs, align=align)
            c.width = Cm(width_cm * col_ratios[ci])


def add_zhengzi_record_table(parent, vocab_list, width_cm=10.5):
    """词汇页底部正字练习记录表 — 8 号字，1 正 + 1 横线，含检查者."""
    add_label(parent, "正字练习记录表", tight=True)
    p = parent.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    add_text(p, "请使用", size=8, color=BLACK)
    add_text(p, "配套字帖练习册", size=8, label=True)
    add_text(p, "进行正字临摹；每练一次，在下方记录一次。", size=8, color=BLACK)

    headers = ["序号", "单词", "正字记录", "练习次数", "检查者"]
    rows = 1 + min(8, len(vocab_list))
    table = parent.add_table(rows=rows, cols=5)
    remove_table_borders(table)
    set_table_width(table, width_cm)
    col_w = [0.7, 2.0, 2.6, 2.2, 2.0]
    for i, h in enumerate(headers):
        c = table.rows[0].cells[i]
        set_cell_borders(c, **learning_frame())
        cell_text(c, h, size=8, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        c.width = Cm(col_w[i])
    set_row_height(table.rows[0], 0.55)

    for ri in range(1, rows):
        row = table.rows[ri]
        set_row_height(row, 0.58)
        if ri <= len(vocab_list):
            word = vocab_list[ri - 1][0]
            vals = [str(ri), word, "正 ___", "□ □ □ □ □", ""]
        else:
            vals = [str(ri), "", "正 ___", "□ □ □ □ □", ""]
        for ci, val in enumerate(vals):
            c = row.cells[ci]
            set_cell_margins(c, top=40, bottom=40, left=50, right=50)
            set_cell_borders(c,
                             bottom={"val": "single", "sz": "4", "color": LEARNING_BLUE},
                             left={"val": "single", "sz": "4", "color": LEARNING_BLUE},
                             right={"val": "single", "sz": "4", "color": LEARNING_BLUE})
            align = WD_ALIGN_PARAGRAPH.LEFT if ci == 1 else WD_ALIGN_PARAGRAPH.CENTER
            cell_text(c, val, size=8, align=align)


def add_en_grid_block(parent, num, en, cn, width_cm=10.5, compact=False, cn_below=False):
    """标准模块：编号+英中文同行（可换行）+ 四线格."""
    specs = FOUR_LINE_SPECS

    wrapper = parent.add_table(rows=5, cols=1)
    remove_table_borders(wrapper)
    set_table_width(wrapper, width_cm)
    for row in wrapper.rows:
        set_row_cant_split(row)

    text_cell = wrapper.rows[0].cells[0]
    set_cell_margins(text_cell, top=50, bottom=80, left=60, right=60)

    text_h = inline_text_height(num, en, cn, chars_per_line=34 if len(en) > 40 else 36)
    set_row_height(wrapper.rows[0], text_h, rule="atLeast")
    add_inline_en_cn(text_cell.paragraphs[0], num, en, cn)

    for i, spec in enumerate(specs):
        apply_grid_line_row(wrapper.rows[i + 1], *spec, width_cm=width_cm)

    add_module_spacer(parent)


def phrase_block_height(en, cn, num=1):
    return inline_text_height(num, en, cn) + FOUR_LINE_GRID_H + MODULE_TAIL_CM + 0.35


def sentence_block_height(en, cn, num=1):
    return inline_text_height(num, en, cn, chars_per_line=34) + FOUR_LINE_GRID_H + MODULE_TAIL_CM + 0.35


class LeftColumnPaginator:
    """内容溢出时自动新开一页，每页均含右侧三模块笔记."""

    def __init__(self, doc, page_num, tip=None, budget=LEFT_COL_BUDGET_CM):
        self.doc = doc
        self.page_num = page_num
        self.tip = tip
        self.budget = budget
        self.left = None
        self.used = 0.0
        self._open = False

    def start(self, main_title=None):
        self._new_page(continue_title=None)
        if main_title:
            add_label(self.left, main_title)
            self.used += 0.45
        return self.left

    def _new_page(self, continue_title=None):
        if self._open:
            finish(self.doc)
        self.left = begin_content_page(self.doc, self.page_num, tip=self.tip)
        self.used = 1.25
        if continue_title:
            add_label(self.left, continue_title, tight=True)
            self.used += 0.42
        self._open = True

    def ensure(self, need_cm, continue_title):
        if self.left is None:
            self._new_page(continue_title)
            return
        if self.used + need_cm > self.budget:
            self._new_page(continue_title)

    def add_cm(self, cm):
        self.used += cm

    def end(self):
        if self._open:
            finish(self.doc)
            self._open = False


def add_phrase_items_paginated(pag, phrases, continue_title, start_num=1, width_cm=10.5):
    if not phrases:
        return
    for i, (en, cn) in enumerate(phrases, start_num):
        h = phrase_block_height(en, cn, num=i)
        pag.ensure(h, continue_title)
        add_en_grid_block(pag.left, i, en, cn, width_cm)
        pag.add_cm(h)


def add_sentence_items_paginated(pag, sentences, continue_title, start_num=1, width_cm=10.5):
    if not sentences:
        return
    for i, (en, cn) in enumerate(sentences, start_num):
        h = sentence_block_height(en, cn, num=i)
        pag.ensure(h, continue_title)
        add_en_grid_block(pag.left, i, en, cn, width_cm)
        pag.add_cm(h)


def add_phrase_section(parent, label, phrases, width_cm=10.5, compact=True):
    if not phrases:
        return
    if label:
        add_label(parent, label, tight=True)
    for i, (en, cn) in enumerate(phrases, 1):
        add_en_grid_block(parent, i, en, cn, width_cm)


def add_full_page_dictation_grids(parent, width_cm=10.5, height_cm=20.5):
    """句型默写专页 — 四线格尽量占满左侧书写区."""
    grid_h = sum(h for h, _, _, _ in FOUR_LINE_SPECS)
    count = max(12, int(height_cm / (grid_h + 0.04)))
    for _ in range(count):
        add_four_line_grid(parent, width_cm)


def add_phrase_items(parent, phrases, width_cm=10.5):
    for i, (en, cn) in enumerate(phrases, 1):
        add_en_grid_block(parent, i, en, cn, width_cm)


def add_sentence_items(parent, sentences, width_cm=10.5):
    for i, (en, cn) in enumerate(sentences, 1):
        add_en_grid_block(parent, i, en, cn, width_cm)


def add_ruled_lines_in_cell(cell, lines=3, line_h=0.28):
    """在单元格内绘制横格线."""
    inner = cell.add_table(rows=lines, cols=1)
    remove_table_borders(inner)
    for row in inner.rows:
        set_row_height(row, line_h)
        c = row.cells[0]
        set_cell_borders(c, bottom={"val": "single", "sz": "4", "color": HEX_LINE})
        c.paragraphs[0].paragraph_format.space_before = Pt(0)
        c.paragraphs[0].paragraph_format.space_after = Pt(0)


def add_error_log_table(parent, width_cm=10.5, count=16):
    """易错整理专页 — 序号 + 横格线书写区."""
    headers = ["序号", "错题内容", "错误原因", "正确写法", "复习□"]
    table = parent.add_table(rows=1 + count, cols=5)
    set_table_width(table, width_cm)
    col_w = [0.7, 4.5, 2.0, 2.8, 0.5]

    for i, h in enumerate(headers):
        c = table.rows[0].cells[i]
        set_cell_borders(c, **learning_frame())
        cell_text(c, h, size=8, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        c.width = Cm(col_w[i])
    set_row_height(table.rows[0], 0.52)

    for ri in range(1, count + 1):
        row = table.rows[ri]
        set_row_height(row, 1.15)
        for ci in range(5):
            c = row.cells[ci]
            set_cell_borders(c,
                             bottom={"val": "single", "sz": "4", "color": LEARNING_BLUE},
                             left={"val": "single", "sz": "4", "color": LEARNING_BLUE},
                             right={"val": "single", "sz": "4", "color": LEARNING_BLUE})
            if ci == 0:
                cell_text(c, str(ri), size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
            elif ci == 4:
                cell_text(c, "□", size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
            elif ci == 1:
                add_ruled_lines_in_cell(c, lines=3, line_h=0.34)
            else:
                add_ruled_lines_in_cell(c, lines=2, line_h=0.34)


def two_column(doc):
    layout = doc.add_table(rows=1, cols=2)
    remove_table_borders(layout)
    set_table_width(layout, 17.4)
    left, right = layout.rows[0].cells[0], layout.rows[0].cells[1]
    left.width = Cm(10.5)
    right.width = Cm(6.9)
    for cell in (left, right):
        cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    return left, right


def begin_content_page(doc, page_num, tip=None):
    """除 Overview 外所有页面：双栏 + 右侧三模块笔记（从页顶开始）."""
    left, right = two_column(doc)
    add_notes_sidebar(right, tip=tip)
    p = left.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    add_text(p, f"Page {page_num:02d}", size=9, label=True)
    add_text(p, f"  ·  {PAGE_LABELS[page_num - 1]}", size=8, label=True)
    return left


def page_with_sidebar(doc, tip=None):
    """兼容旧调用 — 仅建双栏+侧栏（页眉已在外部时使用）."""
    left, right = two_column(doc)
    add_notes_sidebar(right, tip=tip)
    return left


def add_notes_sidebar(cell, tip=None, width_cm=6.5):
    """三模块笔记栏 — 单页 A4 内，横线格书写."""
    add_label(cell, "我的课堂笔记 My Notes", tight=True)
    add_ruled_box(cell, SB_NOTES_H, width_cm, lines=SB_NOTES_LINES)

    add_label(cell, "我听到的重点 Key Points", tight=True)
    add_ruled_box(cell, SB_POINTS_H, width_cm, lines=SB_POINTS_LINES)

    add_label(cell, "我的发现 My Discovery", tight=True)
    add_ruled_box(cell, SB_DISCOVERY_H, width_cm, lines=SB_DISCOVERY_LINES)

    if tip:
        p = cell.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        add_text(p, "学习小贴士：", size=6, label=True)
        add_text(p, tip, size=6, color=BLACK)


def configure_section(section, u=None):
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.5)
    section.footer_distance = Cm(0.5)

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    for r in list(hp.runs):
        r._element.getparent().remove(r._element)
    add_text(hp, "Learning Companion™", size=8, label=True)
    add_text(hp, "  学习陪伴成长手册", size=7, color=BLACK)
    if u:
        add_text(hp, "\t", size=8)
        add_text(hp, f"Unit {u['num']}  {u['title']}（{u['title_cn']}）", size=8, label=True)
        add_text(hp, "\t", size=8)
        add_text(hp, "Date ___/___/___    Done □", size=7, color=GRAY)
        hp.paragraph_format.tab_stops.add_tab_stop(Cm(7.0), WD_TAB_ALIGNMENT.CENTER)
        hp.paragraph_format.tab_stops.add_tab_stop(Cm(14.0), WD_TAB_ALIGNMENT.RIGHT)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    for r in list(fp.runs):
        r._element.getparent().remove(r._element)
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(fp, "Review 复习打卡  ", size=7, label=True)
    add_text(fp, "□ Day1  □ Day3  □ Day7  □ Day15  □ Day30", size=7)
    add_text(fp, "     成长评分  ", size=7, label=True)
    add_text(fp, "☆ ☆ ☆ ☆ ☆", size=8, color=ORANGE)
    if len(footer.paragraphs) < 2:
        footer.add_paragraph()
    fp2 = footer.paragraphs[1]
    for r in list(fp2.runs):
        r._element.getparent().remove(r._element)
    fp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(fp2, OTWSRR, size=6, color=GRAY, italic=True)


def begin_unit(doc, u):
    configure_section(doc.add_section(), u)


def page_title(doc, num, extra=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    add_text(p, f"Page {num:02d}", size=9, label=True)
    add_text(p, f"  ·  {PAGE_LABELS[num - 1]}", size=8, label=True)


def today_goals_left(parent):
    p = parent.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    add_text(p, "今日目标  ", size=7, label=True)
    add_text(p, "□读 □写 □说", size=7)


def finish(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


# ── Pages ────────────────────────────────────────────────────────────

def build_cover(doc):
    for _ in range(5):
        doc.add_paragraph()
    for txt, sz, bold in [
        ("Learning Companion™", 30, True),
        ("小学英语学习成长陪伴手册", 18, False),
        ("人教版 PEP 六年级上册", 14, False),
        ("Unit 1 – 6  ·  V2.0", 14, False),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, txt, size=sz, label=bold, color=BLACK)
    finish(doc)


def build_intro(doc):
    add_label(doc, "使用方法 How to Use")
    for t in ["每天一页，10～15 分钟", "词汇页记单词，短语/句型页练四线格", "除 Overview 外，每页右侧写笔记", "每单元 13 页：含词汇默写、句型默写、易错整理专页"]:
        p = doc.add_paragraph()
        add_text(p, f"• {t}", size=9)
    finish(doc)


def p01_overview(doc, u):
    """Overview — 无三模块侧栏."""
    page_title(doc, 1)
    add_label(doc, "学习目标 Learning Goals")
    for g in u["goals"]:
        p = doc.add_paragraph()
        add_text(p, f"□  {g}", size=9)

    add_label(doc, "学习导航 Learning Map")
    map_box = add_learning_box(doc, 4.5, width_cm=17.0)
    mp = map_box.paragraphs[0]
    add_text(mp, f"★ {u['title']}（{u['title_cn']}）★", size=9, label=True)
    for branch, items in u["learning_map"].items():
        bp = map_box.add_paragraph()
        add_text(bp, f"├─ {branch}：", size=8, label=True)
        add_text(bp, "  ·  ".join(items), size=7, color=BLACK)

    add_label(doc, "思考空间 Think Space")
    add_learning_box(doc, 4.0, width_cm=17.0)

    add_label(doc, "本单元包含 This Unit Includes")
    add_text(doc.add_paragraph(), "  ".join(f"{n}（{d}）" for n, d in UNIT_INCLUDES), size=8)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, MOTIVATIONAL_QUOTES[(u["num"] - 1) % len(MOTIVATIONAL_QUOTES)],
             size=8, italic=True, color=GRAY)
    finish(doc)


def p02_vocabulary(doc, u):
    """词汇页 — Part A 词汇表 + 底部正字记录表."""
    left = begin_content_page(doc, 2, tip=u["learning_tip"])
    today_goals_left(left)
    add_label(left, "① 核心词汇 Key Vocabulary（Part A）")
    add_vocab_table(left, u["part_a_vocab"], width_cm=10.5, compact=True)
    add_zhengzi_record_table(left, u["part_a_vocab"], width_cm=10.5)
    finish(doc)


def p03_phrases(doc, u):
    tip = "短语要会说会用，读准之后再抄写。"
    cont = "② 重点短语 Phrases（续）"
    pag = LeftColumnPaginator(doc, 3, tip=tip)
    pag.start("② 重点短语 Phrases")
    pa = u.get("part_a_phrases", [])
    pb = u.get("part_b_phrases", [])
    add_phrase_items_paginated(pag, pa, cont, start_num=1)
    add_phrase_items_paginated(pag, pb, cont, start_num=len(pa) + 1)
    pag.end()


def p04_sentences(doc, u):
    """句型页 — 英文+中文同行，过长换行，溢出自动分页保留侧栏."""
    tip = "句型先理解意思，再在四线格上抄写。"
    cont = "③ 重点句型 Sentence Patterns（续）"
    pag = LeftColumnPaginator(doc, 4, tip=tip)
    pag.start("③ 重点句型 Sentence Patterns")
    sa = u.get("part_a_sentences", [])
    sb = u.get("part_b_sentences", [])
    add_sentence_items_paginated(pag, sa, cont, start_num=1)
    add_sentence_items_paginated(pag, sb, cont, start_num=len(sa) + 1)
    pag.end()


def p05_vocab_recall(doc, u):
    """词汇默写专页."""
    left = begin_content_page(doc, 5, tip="回忆比重复更重要。先写，再核对。")
    add_label(left, "词汇默写 Vocabulary Recall")
    add_text(left.add_paragraph(), "不看词汇页，凭记忆在四线格上默写单词。", size=8, color=GRAY, italic=True)
    add_full_page_dictation_grids(left, width_cm=10.5, height_cm=20.0)
    finish(doc)


def p06_sentence_dictation(doc, u):
    """句型默写专页 — 中文提示 + 四线格书写区."""
    left = begin_content_page(doc, 6, tip="句型默写前先回忆中文意思，再写英文。")
    add_label(left, "句型默写 Sentence Dictation", tight=True)
    add_text(left.add_paragraph(), "根据中文句意，在四线格上默写完整英文句子。", size=8, color=GRAY, italic=True)
    all_sents = u.get("part_a_sentences", []) + u.get("part_b_sentences", [])
    for i, (en, cn) in enumerate(all_sents, 1):
        p = left.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = TEXT_LINE_SPACING
        add_text(p, f"{i}. ", size=8, label=True)
        add_text(p, cn, size=8, color=BLACK)
    prompt_h = 0.45 + 0.38 * len(all_sents)
    grid_h = max(10.0, 20.5 - prompt_h)
    add_full_page_dictation_grids(left, width_cm=10.5, height_cm=grid_h)
    finish(doc)


def p07_error_log(doc, u):
    """易错整理专页 — 充满横格线与错题序号."""
    left = begin_content_page(doc, 7, tip="错题是最好的老师，写清楚错因才能避免再错。")
    add_label(left, "易错整理 Error Log", tight=True)
    p = left.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    add_text(p, "把做错的单词、短语、句子写进表格，标序号，复习后打勾。", size=7, color=GRAY, italic=True)
    add_error_log_table(left, width_cm=10.5, count=16)
    finish(doc)


def p08_part_b(doc, u):
    tip = u["learning_tip"]
    cont = "⑧ Part B 课堂学习（续）"
    pag = LeftColumnPaginator(doc, 8, tip=tip)
    pag.start("⑧ Part B 课堂学习 Part B Learning")
    today_goals_left(pag.left)
    pag.add_cm(0.35)
    add_label(pag.left, "Part B 核心词汇")
    pag.add_cm(0.45)
    add_vocab_table(pag.left, u["part_b_vocab"], width_cm=10.5)
    pag.add_cm(min(LEFT_COL_BUDGET_CM, 0.5 + 0.48 * len(u["part_b_vocab"])))
    add_phrase_items_paginated(pag, u["part_b_phrases"], cont, start_num=1)
    add_sentence_items_paginated(pag, u["part_b_sentences"], cont, start_num=1)
    pag.end()


def p09_reading(doc, u):
    left = begin_content_page(doc, 9, tip="阅读时不查词典，先猜后验证。")
    add_label(left, f"阅读 {u['reading_title']}")
    box = add_learning_box(left, 3.5, width_cm=10.5)
    add_text(box.paragraphs[0], u["reading_text"], size=9, italic=True, color=DARK)
    add_label(left, "关键词 Key Words")
    add_text(left.add_paragraph(), "提示：" + " · ".join(u["reading_keywords"]), size=8, color=GRAY)
    add_learning_box(left, 2.0, width_cm=10.5)
    add_label(left, "阅读总结 Summary")
    add_learning_box(left, 2.5, width_cm=10.5)
    add_label(left, "阅读思考")
    add_learning_box(left, 2.5, width_cm=10.5)
    finish(doc)


def p10_grammar(doc, u):
    g = u["grammar_discover"]
    left = begin_content_page(doc, 10, tip="语法是自己发现的，记得最牢。")
    add_label(left, f"发现 Discover · {g['title']}")
    add_label(left, "观察 Observe")
    for item in g["observe"]:
        if isinstance(item, tuple):
            add_en_grid_block(left, "", item[0], item[1], width_cm=10.5)
        else:
            p = left.add_paragraph()
            add_text(p, item, size=8, label=True)
    add_label(left, "思考 Think")
    add_learning_box(left, 2.5, width_cm=10.5)
    add_label(left, "老师提示")
    for h in g["hints"]:
        p = left.add_paragraph()
        add_text(p, f"• {h}", size=8)
    add_label(left, "总结 Write")
    add_four_line_grid(left)
    add_four_line_grid(left)
    finish(doc)


def p11_mindmap(doc, u):
    left = begin_content_page(doc, 11, tip="画出你自己的知识连接。")
    add_label(left, f"Unit {u['num']} 知识地图")
    add_text(left.add_paragraph(), f"中心：{u['title']}（{u['title_cn']}）", size=9, label=True)
    add_learning_box(left, 13.0, width_cm=10.5)
    finish(doc)


def p12_challenge(doc, u):
    left = begin_content_page(doc, 12, tip="挑战不是考试，检测学会了多少。")
    add_label(left, "五分钟综合挑战")
    for i, q in enumerate(u["mini_challenge"], 1):
        p = left.add_paragraph()
        add_text(p, f"{i}. {q}", size=9)
        add_four_line_grid(left)
    finish(doc)


def p13_growth(doc, u):
    left = begin_content_page(doc, 13, tip="记录成长，不是记录分数。")
    add_label(left, "成长记录 Growth Journal")
    for lbl in ["今天最大的收获", "今天最大的困难", "下一步目标", "成长感受"]:
        add_label(left, lbl)
        add_learning_box(left, 2.6, width_cm=10.5)
    add_label(left, "复习计划 □Day1 □Day3 □Day7 □Day15 □Day30")
    finish(doc)


def build_unit(doc, u):
    begin_unit(doc, u)
    p01_overview(doc, u)
    p02_vocabulary(doc, u)
    p03_phrases(doc, u)
    p04_sentences(doc, u)
    p05_vocab_recall(doc, u)
    p06_sentence_dictation(doc, u)
    p07_error_log(doc, u)
    p08_part_b(doc, u)
    p09_reading(doc, u)
    p10_grammar(doc, u)
    p11_mindmap(doc, u)
    p12_challenge(doc, u)
    p13_growth(doc, u)


def main():
    doc = Document()
    configure_section(doc.sections[0])
    build_cover(doc)
    build_intro(doc)
    for unit in UNITS:
        build_unit(doc, unit)
    os.makedirs(os.path.dirname(OUTPUT), exist_ok=True)
    doc.save(OUTPUT)
    print(f"Generated: {OUTPUT}")
    print(f"Size: {os.path.getsize(OUTPUT):,} bytes")
    print("Structure: 6 units × 13 pages (Overview无侧栏 + 词汇/句型默写专页 + 易错整理)")


if __name__ == "__main__":
    main()
