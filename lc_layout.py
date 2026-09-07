# -*- coding: utf-8 -*-
"""Layout helpers for Learning Companion V3.0."""
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# ── V3.0 Visual DNA ──────────────────────────────────────────────────
LEARNING_BLUE = "7CB8E8"
LEARNING_BLUE_LIGHT = "B8D9F2"
LEARNING_BLUE_PALE = "E8F4FC"
TITLE_GREEN = "548235"
TIP_ORANGE = "ED7D31"
ERROR_RED = "C00000"
TEXT_DARK = "333333"
TEXT_GRAY = "666666"
WHITE = "FFFFFF"

FONT_CN = "微软雅黑"
FONT_EN = "Calibri"
OTWSRR = "Observe → Think → Write → Speak → Recall → Reflect"

BLUE = RGBColor(0x7C, 0xB8, 0xE8)
GREEN = RGBColor(0x54, 0x82, 0x35)
ORANGE = RGBColor(0xED, 0x7D, 0x31)
RED = RGBColor(0xC0, 0x00, 0x00)
GRAY = RGBColor(0x66, 0x66, 0x66)
DARK = RGBColor(0x33, 0x33, 0x33)

BORDER_SZ = "10"  # ~1.2pt
PAGES_PER_UNIT = 12

FOUR_LINE_SPECS = [
    (0.32, LEARNING_BLUE, "single", "8"),
    (0.30, LEARNING_BLUE_LIGHT, "dashed", "4"),
    (0.30, LEARNING_BLUE_LIGHT, "dashed", "4"),
    (0.32, LEARNING_BLUE, "single", "10"),
]

PAGE_TYPES = [
    "Unit Guide · 学习导航",
    "Vocabulary · 核心词汇",
    "Vocabulary Review · 词汇复习",
    "Phrases · 重点短语",
    "Sentence Patterns · 重点句型",
    "Part A Learning · 课堂学习",
    "Reading Companion · 阅读记录",
    "Grammar Discovery · 语法发现",
    "Knowledge Map · 思维导图",
    "Mini Challenge · 综合挑战",
    "Growth Journal · 成长记录",
    "Review Timeline · 复习时间线",
]

PAGE_REMINDERS = [
    "先浏览整单元地图，建立学习方向，不用一次学完。",
    "今天不用全部记住，先把单词读准、听清。",
    "回忆比重复更重要，先默写，再翻开核对。",
    "短语要会说会用，读准之后再抄写。",
    "句型先理解意思，再模仿造句。",
    "课堂上学到的，当天记下来。",
    "阅读时不查词典，先猜词义，后验证。",
    "语法是自己发现的，观察句子找规律。",
    "思维导图没有标准答案，画出你的理解。",
    "挑战不是考试，检测今天学会了多少。",
    "记录成长，不是记录分数。",
    "按复习时间线打卡，让记忆更长久。",
]


def set_run_font(run, size=10, bold=False, color=None, italic=False):
    run.font.name = FONT_EN
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), FONT_EN)
    rFonts.set(qn("w:hAnsi"), FONT_EN)
    rFonts.set(qn("w:eastAsia"), FONT_CN)
    rPr.insert(0, rFonts)


def add_text(paragraph, text, size=10, bold=False, color=None, italic=False):
    run = paragraph.add_run(text)
    set_run_font(run, size, bold, color, italic)
    return run


def set_cell_shading(cell, color_hex):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color_hex)
    shading.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


def frame_border(spec=None):
    if spec is None:
        spec = {"val": "single", "sz": BORDER_SZ, "color": LEARNING_BLUE}
    return spec


def set_cell_border(cell, edge, val="single", sz=BORDER_SZ, color=LEARNING_BLUE):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    element = OxmlElement(f"w:{edge}")
    element.set(qn("w:val"), val)
    element.set(qn("w:sz"), str(sz))
    element.set(qn("w:space"), "0")
    element.set(qn("w:color"), color)
    tcBorders.append(element)


def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    for edge, spec in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        if spec:
            set_cell_border(cell, edge, spec.get("val", "single"),
                            spec.get("sz", BORDER_SZ), spec.get("color", LEARNING_BLUE))


def set_row_height(row, cm):
    trPr = row._tr.get_or_add_trPr()
    trHeight = OxmlElement("w:trHeight")
    trHeight.set(qn("w:val"), str(int(cm * 567)))
    trHeight.set(qn("w:hRule"), "exact")
    trPr.append(trHeight)


def set_table_width(table, cm):
    tblPr = table._tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        table._tbl.insert(0, tblPr)
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), str(int(cm * 567)))
    tblW.set(qn("w:type"), "dxa")
    tblPr.append(tblW)


def remove_table_borders(table):
    tblPr = table._tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        table._tbl.insert(0, tblPr)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)


def learning_frame_border():
    b = frame_border()
    return dict(top=b, bottom=b, left=b, right=b)


def add_learning_frame(parent, label="", height_cm=2.5, width_cm=17.0):
    """Blue Learning Frame — white interior, free writing (中文思考区)."""
    if label:
        p = parent.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        add_text(p, label, size=8, bold=True, color=GREEN)
    table = parent.add_table(rows=1, cols=1)
    remove_table_borders(table)
    set_table_width(table, width_cm)
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, WHITE)
    set_cell_borders(cell, **learning_frame_border())
    set_cell_margins(cell, 80, 80, 100, 100)
    set_row_height(table.rows[0], height_cm)
    return cell


def add_four_line_grid(parent, width_cm=17.0, compact=False):
    """Publisher blue four-line grid for English writing."""
    table = parent.add_table(rows=4, cols=1)
    remove_table_borders(table)
    set_table_width(table, width_cm)
    specs = FOUR_LINE_SPECS
    if compact:
        specs = [(h * 0.85, c, v, s) for h, c, v, s in specs]
    for i, (h, color, val, sz) in enumerate(specs):
        row = table.rows[i]
        set_row_height(row, h)
        cell = row.cells[0]
        set_cell_borders(cell, bottom={"val": val, "sz": sz, "color": color})
        cell.paragraphs[0].paragraph_format.space_before = Pt(0)
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)


def add_ebbinghaus_row(parent, prefix=""):
    p = parent.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    add_text(p, prefix, size=7, color=GRAY)
    add_text(p, "Day1 □  Day3 □  Day7 □  Day15 □  Day30 □", size=7, color=DARK)


def add_practice_count(parent, word=""):
    p = parent.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    add_text(p, "练习统计  ", size=7, bold=True, color=GREEN)
    if word:
        add_text(p, f"{word}  ", size=7, color=DARK)
    add_text(p, "正 ___  正 ___  丨  ", size=7, color=DARK)
    add_text(p, "次数 □ □ □ □ □", size=7, color=GRAY)


def growth_trail(page_num):
    filled = "🌱" * page_num
    empty = "○" * (PAGES_PER_UNIT - page_num)
    return filled + empty


def configure_section(section, u=None):
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.4)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.5)
    section.footer_distance = Cm(0.5)

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    for r in list(hp.runs):
        r._element.getparent().remove(r._element)
    add_text(hp, "Learning Companion™", size=8, bold=True, color=BLUE)
    add_text(hp, "  学习陪伴成长手册", size=7, color=GRAY)
    if u:
        add_text(hp, "\t", size=8)
        add_text(hp, f"Unit {u['num']}  {u['title']}（{u['title_cn']}）", size=8, bold=True, color=BLUE)
        add_text(hp, "\t", size=8)
        add_text(hp, "Date ___/___/___    Done □", size=7, color=GRAY)
        hp.paragraph_format.tab_stops.add_tab_stop(Cm(7.0), WD_TAB_ALIGNMENT.CENTER)
        hp.paragraph_format.tab_stops.add_tab_stop(Cm(14.0), WD_TAB_ALIGNMENT.RIGHT)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    for r in list(fp.runs):
        r._element.getparent().remove(r._element)
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(fp, "Review 复习打卡  ", size=7, bold=True, color=GREEN)
    add_text(fp, "□ Day1  □ Day3  □ Day7  □ Day15  □ Day30", size=7, color=DARK)
    add_text(fp, "     Today Confidence  ", size=7, bold=True, color=GREEN)
    add_text(fp, "☆ ☆ ☆ ☆ ☆", size=8, color=ORANGE)
    if len(footer.paragraphs) < 2:
        footer.add_paragraph()
    fp2 = footer.paragraphs[1]
    for r in list(fp2.runs):
        r._element.getparent().remove(r._element)
    fp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(fp2, OTWSRR, size=6, color=GRAY, italic=True)


def add_page_chrome(doc, u, page_num):
    """Growth trail, today goals, page title."""
    top = doc.add_table(rows=1, cols=2)
    remove_table_borders(top)
    set_table_width(top, 17.4)
    left, right = top.rows[0].cells[0], top.rows[0].cells[1]

    lp = left.paragraphs[0]
    add_text(lp, f"Page {page_num:02d}", size=10, bold=True, color=BLUE)
    add_text(lp, f"  ·  {PAGE_TYPES[page_num - 1]}", size=9, bold=True, color=GREEN)

    rp = right.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(rp, "🌱 Today  ", size=7, bold=True, color=GREEN)
    add_text(rp, "□读 □写 □说", size=7, color=DARK)

    trail_p = doc.add_paragraph()
    trail_p.paragraph_format.space_before = Pt(2)
    trail_p.paragraph_format.space_after = Pt(4)
    add_text(trail_p, "Growth Trail 成长轨迹  ", size=7, bold=True, color=GREEN)
    add_text(trail_p, growth_trail(page_num), size=8, color=GREEN)
    add_text(trail_p, f"  ({page_num}/{PAGES_PER_UNIT})", size=7, color=GRAY)


def add_module_title(parent, title, level=2):
    p = parent.add_paragraph()
    space = {1: 8, 2: 6, 3: 4}.get(level, 4)
    p.paragraph_format.space_before = Pt(space)
    p.paragraph_format.space_after = Pt(2)
    size = {1: 11, 2: 10, 3: 9}.get(level, 9)
    add_text(p, title, size=size, bold=True, color=GREEN if level <= 2 else BLUE)


def add_learning_reminder(doc, page_num):
    doc.add_paragraph().paragraph_format.space_before = Pt(10)
    cell = add_learning_frame(doc, label="🌱 学习提醒", height_cm=0.9, width_cm=17.0)
    add_text(cell.paragraphs[0], PAGE_REMINDERS[page_num - 1], size=8, color=DARK)


def finish_page(doc, page_num):
    add_learning_reminder(doc, page_num)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_break(WD_BREAK.PAGE)
