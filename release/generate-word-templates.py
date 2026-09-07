#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate editable Word templates using Page N markers (auto-mapped via pdf-page-map)."""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document  # noqa: E402
from textbook_catalog import UNITS  # noqa: E402

OUT_DIR = ROOT / "word-sources"

PAGE_MAP = {
    1: {"unit-start": [9], "unit-next": [10], "part-a": [11, 12], "part-b": [13, 14, 15, 16], "part-c": [17, 18], "reading": [19, 20]},
    2: {"unit-start": [21], "unit-next": [22], "part-a": [23, 24], "part-b": [25, 26, 27, 28], "part-c": [29, 30], "reading": [31, 32]},
    3: {"unit-start": [33], "unit-next": [34], "part-a": [35, 36], "part-b": [37, 38, 39, 40], "part-c": [41, 42], "reading": [43, 44]},
    4: {"unit-start": [45], "unit-next": [46], "part-a": [47, 48], "part-b": [49, 50, 51, 52], "part-c": [53, 54], "reading": [55, 56]},
    5: {"unit-start": [57], "unit-next": [58], "part-a": [59, 60], "part-b": [61, 62, 63, 64], "part-c": [65, 66], "reading": [67, 68]},
    6: {"unit-start": [69], "unit-next": [70], "part-a": [71, 72], "part-b": [73, 74, 75, 76], "part-c": [77, 78], "reading": [79, 80]},
}


def add_page(doc: Document, page: int, pairs: list[tuple[str, str]] | None = None, html: str = "") -> None:
    doc.add_paragraph(f"Page {page}")
    if pairs:
        for en, zh in pairs:
            doc.add_paragraph(f"EN: {en}")
            if zh:
                doc.add_paragraph(f"ZH: {zh}")
    if html:
        for line in html.split("\n"):
            if line.strip():
                doc.add_paragraph(line)
    doc.add_paragraph("")


def split_pairs_evenly(pairs: list[tuple[str, str]], pages: list[int]) -> dict[int, list[tuple[str, str]]]:
    if not pairs or not pages:
        return {p: [] for p in pages}
    per = max(1, (len(pairs) + len(pages) - 1) // len(pages))
    out: dict[int, list[tuple[str, str]]] = {}
    for i, pg in enumerate(pages):
        out[pg] = pairs[i * per : (i + 1) * per]
    return out


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for unit in UNITS:
        u = unit["num"]
        doc = Document()
        doc.add_heading(f"Unit {u} 课文词汇重点句型 双语学习资料", 0)
        doc.add_paragraph(
            "说明：用 Page + 页码 分段（如 Page 11），保存后在网页点「刷新译文」。"
            "系统按页码自动对应左侧课本图，无需手写 Part 名称。"
        )
        pmap = PAGE_MAP.get(u, {})
        for pg in pmap.get("unit-start", []):
            add_page(doc, pg, [(f"Unit {u} {unit.get('title', '')}", unit.get("title_cn", ""))])
        for pg in pmap.get("unit-next", []):
            add_page(doc, pg)
        part_a = split_pairs_evenly(unit.get("part_a_sentences", []), pmap.get("part-a", []))
        for pg, pairs in part_a.items():
            add_page(doc, pg, pairs)
        part_b = split_pairs_evenly(unit.get("part_b_sentences", []), pmap.get("part-b", []))
        for pg, pairs in part_b.items():
            add_page(doc, pg, pairs)
        gd = unit.get("grammar_discover") or {}
        part_c = split_pairs_evenly(gd.get("observe", []), pmap.get("part-c", []))
        for pg, pairs in part_c.items():
            add_page(doc, pg, pairs)
        rt = unit.get("reading_text", "")
        reading_pairs = [(s, "") for s in rt.replace(". ", ".\n").split("\n") if s.strip()] if rt else []
        reading = split_pairs_evenly(reading_pairs, pmap.get("reading", []))
        for pg, pairs in reading.items():
            add_page(doc, pg, pairs)
        for pk in ("part-a", "part-b", "part-c", "reading"):
            pages = pmap.get(pk, [])
            if pages:
                add_page(
                    doc,
                    pages[0],
                    html=f"<h3>Unit {u} {pk} 知识点精讲</h3><p>请在此编写精讲内容（可用 HTML）。</p>",
                )
        out = OUT_DIR / f"Unit{u}_课文词汇重点句型_双语学习资料.docx"
        doc.save(str(out))
        print(f"[OK] {out.name}")
    print(f"[DONE] Templates in {OUT_DIR}")


if __name__ == "__main__":
    main()
