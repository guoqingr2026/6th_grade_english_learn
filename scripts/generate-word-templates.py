#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate editable Word templates using textbook book pages (Page 2, 3, ...)."""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document  # noqa: E402
from textbook_catalog import UNITS  # noqa: E402

OUT_DIR = ROOT / "word-sources"

# 课本印刷页 2–13 → 各 Part 的 PNG 页在 pdf-page-map.json 中配置
BOOK_PAGES = {
    "unit-start": [2],
    "unit-next": [3],
    "part-a": [4, 5],
    "part-b": [6, 7, 8, 9],
    "part-c": [10, 11],
    "reading": [12, 13],
}


def add_page(doc: Document, book_page: int, pairs: list[tuple[str, str]] | None = None) -> None:
    doc.add_paragraph(f"Page {book_page}")
    if pairs:
        for en, zh in pairs:
            if en:
                doc.add_paragraph(en)
            if zh:
                doc.add_paragraph(zh)
    doc.add_paragraph("")


def split_pairs_evenly(pairs: list[tuple[str, str]], pages: list[int]) -> dict[int, list[tuple[str, str]]]:
    if not pairs or not pages:
        return {p: [] for p in pages}
    per = max(1, (len(pairs) + len(pages) - 1) // len(pages))
    return {pg: pairs[i * per : (i + 1) * per] for i, pg in enumerate(pages)}


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for unit in UNITS:
        u = unit["num"]
        doc = Document()
        doc.add_heading(f"Unit {u} 课文词汇重点句型 双语学习资料", 0)
        doc.add_paragraph(
            "说明：Page 后的数字为课本印刷页码（如 Page 2），不是图册文件名。"
            "Unit 1 的 Page 2 对应左侧 page-009.png。保存后点网页「刷新译文」。"
        )
        add_page(doc, BOOK_PAGES["unit-start"][0], [
            (f"Unit {u} {unit.get('title', '')}", unit.get("title_cn", "")),
        ])
        add_page(doc, BOOK_PAGES["unit-next"][0])
        for pg, pairs in split_pairs_evenly(unit.get("part_a_sentences", []), BOOK_PAGES["part-a"]).items():
            add_page(doc, pg, pairs)
        for pg, pairs in split_pairs_evenly(unit.get("part_b_sentences", []), BOOK_PAGES["part-b"]).items():
            add_page(doc, pg, pairs)
        gd = unit.get("grammar_discover") or {}
        for pg, pairs in split_pairs_evenly(gd.get("observe", []), BOOK_PAGES["part-c"]).items():
            add_page(doc, pg, pairs)
        rt = unit.get("reading_text", "")
        reading_pairs = [(s, "") for s in rt.replace(". ", ".\n").split("\n") if s.strip()] if rt else []
        for pg, pairs in split_pairs_evenly(reading_pairs, BOOK_PAGES["reading"]).items():
            add_page(doc, pg, pairs)
        out = OUT_DIR / f"Unit{u}_课文词汇重点句型_双语学习资料.docx"
        doc.save(str(out))
        print(f"[OK] {out.name}")
    print(f"[DONE] Templates in {OUT_DIR} (book pages 2–13)")


if __name__ == "__main__":
    main()
