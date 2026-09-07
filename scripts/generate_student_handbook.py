#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate student tutoring handbook (DOCX + PDF) for the training camp app."""

from __future__ import annotations

import json
import re
import sys
from datetime import date
from html import unescape
from html.parser import HTMLParser
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from appendix_parser import parse_appendix_expressions, parse_appendix_words
from knowledge_u12_sections import FILLBLANK_QA_U12
from knowledge_u36_sections import FILLBLANK_QA_U36
from textbook_catalog import UNITS

DOCS_DIR = ROOT / "docs"
READING_DIR = ROOT / "data" / "study-hub" / "library" / "reading"
READING_BANK = ROOT / "data" / "study-hub" / "reading-bank.json"
KNOWLEDGE_DIR = ROOT / "data" / "study-hub" / "library" / "knowledge"
KNOWLEDGE_BANK = ROOT / "data" / "study-hub" / "knowledge-bank.json"
OUTPUT_DOCX = DOCS_DIR / "学生学习辅导手册.docx"
OUTPUT_PDF = DOCS_DIR / "学生学习辅导手册.pdf"

READING_FILE_RE = re.compile(r"^(?P<prefix>intro|appendix|u(?P<unit>\d+))-p(?P<page>\d+)$", re.I)

FONT_CN = "微软雅黑"
FONT_EN = "Calibri"
BLUE = RGBColor(0x1A, 0x5F, 0xA8)
GRAY = RGBColor(0x55, 0x55, 0x55)

FILLBLANK_ALL = {**FILLBLANK_QA_U12, **FILLBLANK_QA_U36}


def set_run_font(run, size=10, bold=False, color=None, italic=False):
    run.font.name = FONT_EN
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    r_pr = run._element.get_or_add_rPr()
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), FONT_EN)
    r_fonts.set(qn("w:hAnsi"), FONT_EN)
    r_fonts.set(qn("w:eastAsia"), FONT_CN)
    r_pr.insert(0, r_fonts)


def add_para(doc, text, size=10, bold=False, color=None, align=None, space_after=6):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_heading(doc, text, level=1):
    sizes = {1: 16, 2: 14, 3: 12}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, size=sizes.get(level, 11), bold=True, color=BLUE if level <= 2 else None)
    return p


def set_cell_shading(cell, hex_color: str):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), hex_color)
    shading.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading)


def add_table(doc, headers: list[str], rows: list[list[str]], col_widths=None):
    if not rows:
        add_para(doc, "（暂无）", size=9, color=GRAY)
        return
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, size=9, bold=True)
        set_cell_shading(hdr[i], "E8F4FC")
    for ri, row in enumerate(rows, 1):
        for ci, val in enumerate(row):
            cell = table.rows[ri].cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            set_run_font(run, size=8)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()


def strip_html(html: str) -> str:
    text = re.sub(r"<[^>]+>", "", html or "")
    return unescape(text).strip()


def inline_from_html(fragment: str) -> str:
    text = fragment or ""
    text = re.sub(r"<br\s*/?>", "\n", text, flags=re.I)
    text = re.sub(r"</(p|div|li|h[1-4])>", "\n", text, flags=re.I)
    return strip_html(text)


def parse_html_table(table_html: str) -> tuple[list[str], list[list[str]]]:
    rows: list[list[str]] = []
    for row_html in re.findall(r"<tr[^>]*>(.*?)</tr>", table_html, re.DOTALL | re.I):
        cells = [
            inline_from_html(cell)
            for cell in re.findall(r"<t[hd][^>]*>(.*?)</t[hd]>", row_html, re.DOTALL | re.I)
        ]
        if any(cells):
            rows.append(cells)
    if not rows:
        return [], []
    max_cols = max(len(r) for r in rows)
    norm = [r + [""] * (max_cols - len(r)) for r in rows]
    return norm[0], norm[1:]


class KnowledgeContentParser(HTMLParser):
    """Convert knowledge lecture HTML into ordered blocks for Word export."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.blocks: list[tuple[str, object]] = []
        self._buf: list[str] = []
        self._block_tag: str | None = None
        self._in_table = False
        self._table_html: list[str] = []
        self._table_depth = 0
        self._in_list = False
        self._list_ordered = False
        self._list_items: list[str] = []
        self._li_buf: list[str] = []
        self._in_li = False

    def _take_buf(self) -> str:
        text = "".join(self._buf).strip()
        self._buf = []
        return text

    def handle_starttag(self, tag, attrs):
        tag = tag.lower()
        if tag in ("script", "style"):
            return
        if tag == "hr":
            self.blocks.append(("hr", None))
            return
        if tag == "table":
            self._in_table = True
            self._table_depth = 1
            self._table_html = [self.get_starttag_text() or "<table>"]
            return
        if self._in_table:
            self._table_html.append(self.get_starttag_text() or f"<{tag}>")
            if tag == "table":
                self._table_depth += 1
            return
        if tag == "ul":
            self._in_list = True
            self._list_ordered = False
            self._list_items = []
            return
        if tag == "ol":
            self._in_list = True
            self._list_ordered = True
            self._list_items = []
            return
        if tag == "li" and self._in_list:
            self._in_li = True
            self._li_buf = []
            return
        if tag in ("h1", "h2", "h3", "h4", "p", "pre", "blockquote"):
            self._block_tag = tag
            self._buf = []
            return
        if tag == "br":
            self._buf.append("\n")
        elif tag == "code" and self._block_tag == "pre":
            pass

    def handle_endtag(self, tag):
        tag = tag.lower()
        if self._in_table:
            self._table_html.append(f"</{tag}>")
            if tag == "table":
                self._table_depth -= 1
                if self._table_depth <= 0:
                    self.blocks.append(("table", "".join(self._table_html)))
                    self._in_table = False
            return
        if tag == "li" and self._in_li:
            item = inline_from_html("".join(self._li_buf))
            if item:
                self._list_items.append(item)
            self._in_li = False
            return
        if tag in ("ul", "ol") and self._in_list:
            self.blocks.append(("list", (self._list_ordered, self._list_items[:])))
            self._in_list = False
            return
        if tag == self._block_tag:
            text = self._take_buf()
            if text:
                self.blocks.append((self._block_tag, text))
            self._block_tag = None
            return

    def handle_data(self, data):
        if self._in_table:
            self._table_html.append(data)
        elif self._in_li:
            self._li_buf.append(data)
        elif self._block_tag or not self._in_list:
            self._buf.append(data)


def parse_knowledge_html(html: str) -> list[tuple[str, object]]:
    cleaned = re.sub(r"<script[^>]*>.*?</script>", "", html or "", flags=re.DOTALL | re.I)
    cleaned = re.sub(r"<style[^>]*>.*?</style>", "", cleaned, flags=re.DOTALL | re.I)
    parser = KnowledgeContentParser()
    parser.feed(cleaned)
    parser.close()
    return parser.blocks


def render_knowledge_blocks(doc, blocks: list[tuple[str, object]]):
    for kind, payload in blocks:
        if kind == "hr":
            continue
        if kind in ("h1", "h2"):
            add_heading(doc, str(payload), 2)
        elif kind == "h3":
            add_heading(doc, str(payload), 3)
        elif kind == "h4":
            add_para(doc, str(payload), size=10, bold=True)
        elif kind == "pre":
            add_para(doc, str(payload), size=8)
        elif kind in ("p", "blockquote"):
            add_para(doc, str(payload), size=9)
        elif kind == "list":
            ordered, items = payload  # type: ignore[misc]
            for i, item in enumerate(items, 1):
                prefix = f"{i}. " if ordered else "• "
                add_para(doc, f"{prefix}{item}", size=9)
        elif kind == "table":
            headers, rows = parse_html_table(str(payload))
            if headers:
                n = len(headers)
                widths = [16.0 / n] * n if n else None
                add_table(doc, headers, rows, col_widths=widths)


def load_knowledge_bank() -> dict:
    if not KNOWLEDGE_BANK.is_file():
        return {}
    return json.loads(KNOWLEDGE_BANK.read_text(encoding="utf-8"))


def list_knowledge_pages(bank: dict | None = None) -> list[dict]:
    bank = bank or load_knowledge_bank()
    pages: list[dict] = []
    for ublock in bank.get("units", []):
        unit = int(ublock.get("unit", 0))
        for page in ublock.get("pages", []):
            pages.append(
                {
                    "id": page.get("id", ""),
                    "label": page.get("label", ""),
                    "unit": unit,
                    "unit_title": ublock.get("title", ""),
                    "section": "unit",
                }
            )
    for page in bank.get("appendix", {}).get("pages", []):
        pages.append(
            {
                "id": page.get("id", ""),
                "label": page.get("label", ""),
                "unit": 0,
                "unit_title": bank.get("appendix", {}).get("title", "附录"),
                "section": "appendix",
            }
        )
    return pages


def knowledge_body_html(page_id: str, bank: dict) -> str:
    lib_path = KNOWLEDGE_DIR / f"{page_id}.html"
    if lib_path.is_file():
        return lib_path.read_text(encoding="utf-8")

    def scan(pages: list[dict]) -> str:
        for page in pages:
            if page.get("id") == page_id:
                return page.get("bodyHtml") or ""
        return ""

    for ublock in bank.get("units", []):
        html = scan(ublock.get("pages", []))
        if html:
            return html
    return scan(bank.get("appendix", {}).get("pages", []))


def knowledge_section_banner(page: dict) -> str | None:
    if page.get("section") == "appendix":
        return "【教材附录 · 知识点精讲】"
    unit = page.get("unit")
    if unit:
        return f"【Unit {unit} · 知识点精讲】"
    return None


def build_knowledge_full_text(doc):
    bank = load_knowledge_bank()
    pages = list_knowledge_pages(bank)
    add_heading(doc, f"第四篇  知识点精讲全文（共 {len(pages)} 页）", 1)
    add_para(
        doc,
        "以下内容与软件「知识点精讲」一致：优先使用 library/knowledge 中的 HTML，"
        "与软件左侧课本图对照复习。按 Unit 1 → Unit 6 → 附录顺序排列。",
        size=9,
    )

    last_banner = None
    for idx, page in enumerate(pages, 1):
        banner = knowledge_section_banner(page)
        if banner and banner != last_banner:
            add_heading(doc, banner, 2)
            last_banner = banner

        title = page.get("label") or page.get("id", "")
        if page.get("unit"):
            title = f"Unit {page['unit']} · {title}"
        add_heading(doc, f"第 {idx} 页 · {title}", 3)

        html = knowledge_body_html(page["id"], bank)
        if not html.strip():
            add_para(doc, "（本页内容为空，请在软件中编辑后重新生成手册。）", size=8, color=GRAY)
            continue

        blocks = parse_knowledge_html(html)
        if not blocks:
            add_para(doc, inline_from_html(html), size=9)
        else:
            render_knowledge_blocks(doc, blocks)

    doc.add_page_break()


def reading_file_sort_key(path: Path) -> tuple:
    m = READING_FILE_RE.match(path.stem)
    if not m:
        return (99, 0, path.name)
    prefix = m.group("prefix").lower()
    page = int(m.group("page"))
    if prefix == "intro":
        return (0, page, path.name)
    if prefix == "appendix":
        return (8, page, path.name)
    return (int(m.group("unit")), page, path.name)


def list_reading_html_files() -> list[Path]:
    files = sorted(READING_DIR.glob("*.html"), key=reading_file_sort_key)
    return files


def parse_reading_html(path: Path) -> list[tuple[str, str]]:
    raw = path.read_text(encoding="utf-8")
    parts = re.findall(r'<p class="(en|zh)"[^>]*>(.*?)</p>', raw, re.DOTALL | re.IGNORECASE)
    pairs: list[tuple[str, str]] = []
    i = 0
    while i < len(parts):
        kind, body = parts[i]
        if kind.lower() == "en":
            en = strip_html(body)
            zh = ""
            if i + 1 < len(parts) and parts[i + 1][0].lower() == "zh":
                zh = strip_html(parts[i + 1][1])
                i += 2
            else:
                i += 1
            if en or zh:
                pairs.append((en, zh))
        elif kind.lower() == "zh":
            zh = strip_html(body)
            if zh:
                pairs.append(("", zh))
            i += 1
        else:
            i += 1
    return pairs


def build_reading_label_lookup() -> dict[str, dict]:
    if not READING_BANK.is_file():
        return {}
    bank = json.loads(READING_BANK.read_text(encoding="utf-8"))
    lookup: dict[str, dict] = {}

    def put(key: str, atlas: int, book, unit=None, section: str = ""):
        lookup[key] = {"atlas": atlas, "book": book, "unit": unit, "section": section}

    for page in bank.get("intro", {}).get("pages", []):
        for pc in page.get("pageContents", []) or []:
            atlas = int(pc.get("page", 0))
            put(f"intro-p{atlas:03d}", atlas, pc.get("bookPage"), section="intro")

    for page in bank.get("appendix", {}).get("pages", []):
        for pc in page.get("pageContents", []) or []:
            atlas = int(pc.get("page", 0))
            put(f"appendix-p{atlas:03d}", atlas, pc.get("bookPage"), section="appendix")

    for ublock in bank.get("units", []):
        unit = int(ublock.get("unit", 0))
        for page in ublock.get("pages", []):
            for pc in page.get("pageContents", []) or []:
                atlas = int(pc.get("page", 0))
                put(f"u{unit}-p{atlas:03d}", atlas, pc.get("bookPage"), unit=unit, section="unit")

    return lookup


def format_reading_page_title(stem: str, lookup: dict[str, dict]) -> str:
    meta = lookup.get(stem, {})
    m = READING_FILE_RE.match(stem)
    atlas = meta.get("atlas") or (int(m.group("page")) if m else 0)
    book = meta.get("book")
    if stem.startswith("intro"):
        return f"开篇 · 图册 page-{atlas:03d}"
    if stem.startswith("appendix"):
        return f"附录 · 图册 page-{atlas:03d}"
    unit = meta.get("unit") or (int(m.group("unit")) if m and m.group("unit") else "?")
    book_txt = f" · 课本 Page {book}" if book else ""
    return f"Unit {unit} · 图册 page-{atlas:03d}{book_txt}"


def reading_section_banner(stem: str) -> str | None:
    if stem.startswith("intro"):
        return "【开篇与目录】"
    if stem.startswith("appendix"):
        return "【教材附录】"
    m = READING_FILE_RE.match(stem)
    if m and m.group("unit"):
        return f"【Unit {int(m.group('unit'))}】"
    return None


def build_reading_full_text(doc):
    files = list_reading_html_files()
    lookup = build_reading_label_lookup()
    add_heading(doc, f"第三篇  课文中英文讲读全文（共 {len(files)} 页）", 1)
    add_para(
        doc,
        "以下内容与软件「课文中英文讲读」一致，按图册页码顺序排列（intro → Unit 1–6 → 附录）。"
        "每页先列英文，对应中文在右侧，便于离线复习与打印朗读。",
        size=9,
    )

    last_banner = None
    page_no = 0
    for path in files:
        page_no += 1
        stem = path.stem
        banner = reading_section_banner(stem)
        if banner and banner != last_banner:
            add_heading(doc, banner, 2)
            last_banner = banner

        title = format_reading_page_title(stem, lookup)
        add_heading(doc, f"第 {page_no} 页 · {title}", 3)
        pairs = parse_reading_html(path)
        if not pairs:
            add_para(doc, "（本页以课本图片为主，见软件左侧课本图。）", size=8, color=GRAY)
        else:
            rows = [[en, zh] for en, zh in pairs]
            add_table(doc, ["英文 English", "中文"], rows, col_widths=[8.5, 7.5])

    doc.add_page_break()


def build_software_guide(doc):
    add_heading(doc, "第一篇  软件使用辅导（学生版）", 1)
    add_para(
        doc,
        "本手册配套「PEP 六年级英语上册综合知识记忆训练营」网页软件使用。"
        "请通过老师或家长提供的局域网地址访问（如 http://192.168.x.x:8080），"
        "不要直接双击 HTML 文件打开。",
        size=9,
    )

    add_heading(doc, "1. 第一次怎么用？", 2)
    steps = [
        "双击电脑上的「启动.bat」（由老师/家长操作），浏览器打开练习页面。",
        "在右侧输入你的昵称并保存（手机与电脑要用相同昵称才能同步）。",
        "主界面顶部选择模块：单词卡片、语法句型、选择题等。",
        "需要跟读课文时，点「课文中英文讲读」；需要老师讲法时，点「知识点精讲」。",
        "做完题想保存进度，点右侧「同步数据」（约 10 秒自动同步一次，也可手动点）。",
    ]
    for i, s in enumerate(steps, 1):
        add_para(doc, f"{i}. {s}", size=9)

    add_heading(doc, "2. 各模块怎么练？", 2)
    modules = [
        ("单词卡片", "看英文想中文，点卡片翻面；不会可请家长判错记入错题本。"),
        ("语法句型", "先看规则，再做「主谓宾连连看」；手动点下一组，不会自动跳题。"),
        ("常用语", "遮盖关键词回忆完整句子；可筛选单元。"),
        ("选择题", "选答案看解析；须自己点「下一题」，答对/答错不会自动换题。"),
        ("综合填空", "每组 5 题，填英文单词；点「下一组」换题。"),
        ("不规则动词", "附录5 共 47 个动词，练原形/过去式/过去分词。"),
        ("优秀作文", "读范文、做填空、可用评分器自评。"),
        ("错题本", "自动收录答错题；可筛选模块与单元复习。"),
        ("学习日志", "用英文写今日感想（建议 100 词以内）。"),
    ]
    add_table(doc, ["模块", "学生操作要点"], modules, col_widths=[3.5, 13.5])

    add_heading(doc, "3. 讲读与精讲怎么用？", 2)
    tips = [
        "打开后左侧是课本图，右侧是英文与中文；不会自动朗读。",
        "只有点工具栏「朗读」才会逐句向下读，读完本页停止，不会自动翻 Part。",
        "精讲含词汇、语法、填空问答与写作范文，建议先讲读再精讲再做题。",
        "可在精讲页写「我的交互笔记」，记得之后点「同步数据」备份。",
    ]
    for t in tips:
        add_para(doc, f"• {t}", size=9)

    add_heading(doc, "4. 积分与错题（学生须知）", 2)
    add_para(doc, "选择题、填空、动词等系统自动判题，首次答对可获得做题积分。", size=9)
    add_para(doc, "同一道题以前答对过，再次答对不再加分，但仍可练习。", size=9)
    add_para(doc, "答错的题会自动进入错题本，建议每周复习一次。", size=9)
    add_para(doc, "家长密码、行为积分、兑换等由家长操作，学生无需知道密码。", size=9, color=GRAY)

    add_heading(doc, "5. 推荐每日 20 分钟学习路线", 2)
    route = [
        "5 分钟：课文中英文讲读（当前 Part）",
        "5 分钟：知识点精讲（看词汇+语法要点）",
        "5 分钟：单词卡片 + 语法/选择题",
        "3 分钟：错题本复习",
        "2 分钟：学习日志（英文写一句收获）",
    ]
    for i, r in enumerate(route, 1):
        add_para(doc, f"步骤 {i}：{r}", size=9)

    doc.add_page_break()


def build_unit_section(doc, unit: dict, appendix_words: dict, appendix_exprs: dict):
    u = unit["num"]
    add_heading(doc, f"Unit {u}  {unit['title']}（{unit['title_cn']}）", 1)
    add_para(doc, f"主题：{unit['theme']}", size=9, color=GRAY)

    add_heading(doc, "学习目标", 2)
    for g in unit.get("goals", []):
        add_para(doc, f"□ {g}", size=9)

    add_heading(doc, "核心词汇 Part A", 2)
    rows = [[w[0], w[1], w[2], w[3]] for w in unit.get("part_a_vocab", [])]
    add_table(doc, ["单词", "音标", "词性", "中文"], rows, col_widths=[4, 3, 1.5, 5.5])

    add_heading(doc, "核心词汇 Part B", 2)
    rows = [[w[0], w[1], w[2], w[3]] for w in unit.get("part_b_vocab", [])]
    add_table(doc, ["单词", "音标", "词性", "中文"], rows, col_widths=[4, 3, 1.5, 5.5])

    add_heading(doc, "常用短语", 2)
    phrases = unit.get("part_a_phrases", []) + unit.get("part_b_phrases", [])
    add_table(doc, ["英文", "中文"], phrases, col_widths=[8, 8])

    add_heading(doc, "重点句型", 2)
    sents = unit.get("part_a_sentences", []) + unit.get("part_b_sentences", [])
    add_table(doc, ["英文", "中文"], sents, col_widths=[8, 8])

    gd = unit.get("grammar_discover", {})
    if gd:
        add_heading(doc, f"语法发现：{gd.get('title', '')}", 2)
        for item in gd.get("observe", []):
            if isinstance(item, tuple):
                add_para(doc, f"• {item[0]}  —  {item[1]}", size=9)
            else:
                add_para(doc, f"• {item}", size=9)
        add_para(doc, "老师提示：", size=9, bold=True)
        for h in gd.get("hints", []):
            add_para(doc, f"  - {h}", size=9)

    add_heading(doc, f"阅读：{unit.get('reading_title', 'Reading')}", 2)
    add_para(doc, unit.get("reading_text", ""), size=9)
    kw = unit.get("reading_keywords", [])
    if kw:
        add_para(doc, "关键词：" + " · ".join(kw), size=9, color=GRAY)

    aw = appendix_words.get(u, [])
    if aw:
        add_heading(doc, f"附录2 本单元词汇（{len(aw)} 词）", 2)
        rows = [[w.display_word, f"/{w.ipa}/", w.gloss, w.pos_page] for w in aw[:40]]
        add_table(doc, ["单词", "音标", "释义", "课本页"], rows, col_widths=[4, 3, 6, 2])
        if len(aw) > 40:
            add_para(doc, f"… 其余 {len(aw) - 40} 词见软件「附录校验」或课本附录2。", size=8, color=GRAY)

    ae = appendix_exprs.get(u, [])
    if ae:
        add_heading(doc, f"附录4 常用表达（{len(ae)} 条）", 2)
        rows = [[e.en, e.zh] for e in ae[:25]]
        add_table(doc, ["英文", "中文"], rows, col_widths=[8, 8])
        if len(ae) > 25:
            add_para(doc, f"… 其余 {len(ae) - 25} 条见软件精讲/课本附录4。", size=8, color=GRAY)

    add_heading(doc, "课文填空与问答（精讲要点）", 2)
    for part_key, label in [("part-a", "Part A"), ("part-b", "Part B"), ("part-c", "Part C"), ("reading", "Reading")]:
        items = FILLBLANK_ALL.get((u, part_key), [])
        if not items:
            continue
        add_para(doc, label, size=10, bold=True)
        for q, a in items:
            add_para(doc, f"问：{q}", size=9)
            add_para(doc, f"答：{a}", size=9, color=GRAY)

    add_heading(doc, "综合小挑战", 2)
    for i, q in enumerate(unit.get("mini_challenge", []), 1):
        add_para(doc, f"{i}. {q}", size=9)

    add_para(doc, f"学习小贴士：{unit.get('learning_tip', '')}", size=9, color=GRAY)
    doc.add_page_break()


def build_irregular_verbs(doc):
    add_heading(doc, "第五篇  附录5 不规则动词表（47 个）", 1)
    add_para(doc, "软件「动词专项」按本表练习三态。请重点记忆过去式与过去分词。", size=9)
    path = ROOT / "data" / "practice" / "irregular-verbs.json"
    data = json.loads(path.read_text(encoding="utf-8"))
    rows = [
        [v["base"], v["past"], v["participle"], v.get("zh", "")]
        for v in data.get("verbs", [])
    ]
    add_table(doc, ["原形", "过去式", "过去分词", "中文"], rows, col_widths=[4, 3.5, 3.5, 5])


def build_appendix_reference(doc, appendix_words, appendix_exprs):
    add_heading(doc, "第六篇  全册附录速查", 1)
    total_w = sum(len(v) for v in appendix_words.values())
    total_e = sum(len(v) for v in appendix_exprs.values())
    add_para(doc, f"附录2 词汇共 {total_w} 条；附录4 表达共 {total_e} 条（与教材一致）。", size=9)
    for u in range(1, 7):
        wc = len(appendix_words.get(u, []))
        ec = len(appendix_exprs.get(u, []))
        add_para(doc, f"Unit {u}：词汇 {wc} 条 · 表达 {ec} 条", size=9)


def build_cover(doc):
    for _ in range(4):
        doc.add_paragraph()
    add_para(
        doc,
        "PEP 六年级英语上册\n综合知识记忆训练营",
        size=22,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=12,
    )
    add_para(
        doc,
        "学生学习辅导手册\n（软件使用 + 全册学习资料 + 112 页课文双语 + 知识点精讲全文）",
        size=16,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=8,
    )
    add_para(
        doc,
        "2026 秋人教版修订版  ·  Unit 1–6",
        size=12,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        color=GRAY,
    )
    add_para(
        doc,
        f"生成日期：{date.today().isoformat()}",
        size=10,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        color=GRAY,
    )
    add_para(
        doc,
        "旭日长空光照人生（Sunshine Life Team）",
        size=10,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        color=GRAY,
    )
    doc.add_page_break()


def build_toc(doc):
    add_heading(doc, "目  录", 1)
    items = [
        "第一篇  软件使用辅导（学生版）",
        "第二篇  Unit 1–6 学习资料（词汇·短语·句型·语法·阅读·精讲问答）",
        "第三篇  课文中英文讲读全文（112 页逐句对照）",
        "第四篇  知识点精讲全文（与软件精讲一致）",
        "第五篇  附录5 不规则动词表（47 个）",
        "第六篇  全册附录速查",
    ]
    for item in items:
        add_para(doc, item, size=10)
    add_para(
        doc,
        "说明：第三篇与「课文中英文讲读」一致；第四篇与「知识点精讲」一致，可离线打印复习。",
        size=9,
        color=GRAY,
    )
    doc.add_page_break()


def configure_doc(doc):
    sec = doc.sections[0]
    sec.page_height = Cm(29.7)
    sec.page_width = Cm(21.0)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)


def export_pdf_from_docx(docx_path: Path, pdf_path: Path) -> bool:
    """Try Word COM (Windows) then pymupdf plain export."""
    try:
        import win32com.client  # type: ignore

        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        doc = word.Documents.Open(str(docx_path.resolve()))
        doc.SaveAs(str(pdf_path.resolve()), FileFormat=17)
        doc.Close()
        word.Quit()
        return pdf_path.is_file()
    except Exception:
        pass

    try:
        import fitz  # pymupdf

        data = json.loads((ROOT / "data" / "practice" / "irregular-verbs.json").read_text(encoding="utf-8"))
        font_paths = [
            Path(r"C:\Windows\Fonts\msyh.ttc"),
            Path(r"C:\Windows\Fonts\simhei.ttf"),
            Path(r"C:\Windows\Fonts\simsun.ttc"),
        ]
        fontfile = next((p for p in font_paths if p.is_file()), None)
        if not fontfile:
            return False

        pdf = fitz.open()
        page = pdf.new_page(width=595, height=842)
        y = 50
        fontname = "china-font"
        page.insert_font(fontname=fontname, fontfile=str(fontfile))

        def writeln(text, size=11, gap=16):
            nonlocal y, page
            for line in text.split("\n"):
                if y > 780:
                    page = pdf.new_page(width=595, height=842)
                    page.insert_font(fontname=fontname, fontfile=str(fontfile))
                    y = 50
                page.insert_text((50, y), line, fontname=fontname, fontsize=size)
                y += gap

        writeln("PEP 六年级英语上册 · 学生学习辅导手册（PDF 简版）", 14)
        writeln("")
        writeln("完整版请打开同目录 Word 文件：学生学习辅导手册.docx", 10)
        writeln("")
        writeln("软件访问：启动.bat 后浏览器打开 http://localhost:8080", 10)
        writeln("")
        writeln("附录5 不规则动词（47）：", 12)
        for v in data.get("verbs", []):
            writeln(f"{v['base']}  →  {v['past']}  /  {v['participle']}  ({v.get('zh','')})", 9, 13)

        pdf.save(str(pdf_path))
        pdf.close()
        return True
    except Exception as exc:
        print(f"PDF export fallback failed: {exc}")
        return False


def generate_docx() -> Path:
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    appendix_words = parse_appendix_words()
    appendix_exprs = parse_appendix_expressions()

    doc = Document()
    configure_doc(doc)
    build_cover(doc)
    build_toc(doc)
    build_software_guide(doc)

    add_heading(doc, "第二篇  Unit 1–6 学习资料", 1)
    doc.add_page_break()
    for unit in UNITS:
        build_unit_section(doc, unit, appendix_words, appendix_exprs)

    build_reading_full_text(doc)
    build_knowledge_full_text(doc)
    build_irregular_verbs(doc)
    build_appendix_reference(doc, appendix_words, appendix_exprs)

    add_heading(doc, "结语", 1)
    add_para(
        doc,
        "祝学习进步！请坚持每天使用软件讲读、精讲与练习，"
        "错题本每周复习，期末对照本手册与课本附录查漏补缺。",
        size=10,
    )
    add_para(doc, "— Sunshine Life Team", size=10, align=WD_ALIGN_PARAGRAPH.RIGHT)

    doc.save(str(OUTPUT_DOCX))
    return OUTPUT_DOCX


def main():
    print("=" * 60)
    print("Student Handbook Generator")
    print("=" * 60)
    docx_path = generate_docx()
    print(f"[OK] DOCX -> {docx_path} ({docx_path.stat().st_size:,} bytes)")
    reading_count = len(list_reading_html_files())
    knowledge_count = len(list_knowledge_pages())
    print(f"     课文双语页数: {reading_count}")
    print(f"     知识点精讲页数: {knowledge_count}")

    if export_pdf_from_docx(docx_path, OUTPUT_PDF):
        print(f"[OK] PDF  -> {OUTPUT_PDF} ({OUTPUT_PDF.stat().st_size:,} bytes)")
    else:
        print("[WARN] PDF not generated. Install Microsoft Word or use DOCX print-to-PDF.")

    print()
    print("Students: open docs/学生学习辅导手册.docx or print to PDF from Word.")


if __name__ == "__main__":
    main()
